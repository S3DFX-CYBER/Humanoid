"""End-to-end pipeline test: research → outline → draft → verify → style → format.

Search, the LLM, and the render service are stubbed, but the database, the
stage transitions, the citation numbering, and document persistence are real.
"""

import io
import json
import re
import uuid

from tests.conftest import requires_db

pytestmark = requires_db

SOURCE_ID_RE = re.compile(r"SOURCE_ID: ([0-9a-f-]{36})")
STYLE_TEXT_RE = re.compile(r"TEXT:\n(.*)\n\nReturn only the rewritten text\.", re.S)


class FakeProviderPool:
    """Deterministic stand-in for the provider pool."""

    def __init__(self):
        self.prompts: list[str] = []

    async def call(self, prompt: str, tier: str = "cheap", **kwargs) -> str:
        self.prompts.append(prompt)

        if "fact checker" in prompt:
            return "pass"

        if "precise academic editor" in prompt:
            match = STYLE_TEXT_RE.search(prompt)
            assert match, "style prompt did not carry the section text"
            return match.group(1)

        if "drafting one section" in prompt:
            match = SOURCE_ID_RE.search(prompt)
            assert match, "draft prompt carried no source id"
            source_id = match.group(1)
            return (
                f"Transformers removed recurrence [cite: {source_id}]. "
                f"Long sequences improved as a result [cite: {source_id}]."
            )

        raise AssertionError(f"unexpected prompt: {prompt[:120]}")

    async def call_json(self, prompt: str, tier: str = "cheap", **kwargs) -> str:
        self.prompts.append(prompt)

        if "research librarian" in prompt:
            return json.dumps({"queries": ["transformer architectures"]})

        if "planning a research document" in prompt:
            return json.dumps(
                {
                    "title": "Transformers in Practice",
                    "thesis": "Attention replaced recurrence.",
                    "sections": ["Introduction", "Conclusion"],
                }
            )

        raise AssertionError(f"unexpected json prompt: {prompt[:120]}")


class FakeResponse:
    def __init__(self, content: bytes):
        self.content = content

    def raise_for_status(self) -> None:
        return None


class FakeHttpClient:
    """Stands in for the render service, returning a real DOCX."""

    def __init__(self, *args, **kwargs):
        pass

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc_info):
        return False

    async def post(self, url, json=None):  # noqa: A002 - mirrors httpx
        """Return a small but real DOCX, independent of the render service."""
        from docx import Document

        document = Document()
        document.add_heading(json["title"], level=1)
        for section in json["sections"]:
            document.add_heading(section["header"], level=2)
            document.add_paragraph(section["body"])
        document.add_heading("References", level=1)
        for index, reference in enumerate(json["references"], 1):
            document.add_paragraph(f"[{index}] {reference['title']}")

        buffer = io.BytesIO()
        document.save(buffer)
        return FakeResponse(buffer.getvalue())


async def test_pipeline_runs_end_to_end(db_pool, monkeypatch):
    from api.database import ensure_user
    from workers import stages
    from workers.worker import run_pipeline

    user_id = str(uuid.uuid4())
    job_id = str(uuid.uuid4())
    await ensure_user(user_id)

    async with db_pool.acquire() as conn:
        await conn.execute(
            """
            INSERT INTO jobs (id, user_id, topic, status, citation_style)
            VALUES ($1, $2, 'Transformers in machine translation', 'queued', 'apa')
            """,
            uuid.UUID(job_id),
            uuid.UUID(user_id),
        )

    async def fake_search(query, max_results=3):
        return [
            {
                "title": "Attention Is All You Need",
                "href": "https://example.org/attention",
                "body": "A" * 300,
            }
        ]

    async def fake_fetch(url, max_chars=8000):
        return (
            "Attention mechanisms replaced recurrence entirely, which reduced "
            "training time and improved translation quality on long sequences. "
        ) * 4

    async def fake_embedding(text):
        return None

    monkeypatch.setattr(stages, "get_search_results", fake_search)
    monkeypatch.setattr(stages, "fetch_and_extract_text", fake_fetch)
    monkeypatch.setattr(stages, "generate_embedding", fake_embedding)
    monkeypatch.setattr(stages.httpx, "AsyncClient", FakeHttpClient)

    provider = FakeProviderPool()
    result = await run_pipeline(
        {"db_pool": db_pool, "provider_pool": provider},
        job_id,
        user_id,
    )
    assert result == "done"

    async with db_pool.acquire() as conn:
        status = await conn.fetchval(
            "SELECT status::text FROM jobs WHERE id = $1", uuid.UUID(job_id)
        )
        stage_rows = await conn.fetch(
            """
            SELECT stage_name, status::text AS status FROM job_stages
            WHERE job_id = $1 ORDER BY started_at
            """,
            uuid.UUID(job_id),
        )
        source_count = await conn.fetchval(
            "SELECT count(*) FROM sources WHERE job_id = $1", uuid.UUID(job_id)
        )
        document = await conn.fetchrow(
            "SELECT filename, byte_size, content FROM documents WHERE job_id = $1",
            uuid.UUID(job_id),
        )
        format_output = await conn.fetchval(
            """
            SELECT output_data FROM job_stages
            WHERE job_id = $1 AND stage_name = 'format'
            """,
            uuid.UUID(job_id),
        )

    assert status == "done"
    assert [row["stage_name"] for row in stage_rows] == [
        "research",
        "outline",
        "draft",
        "verify",
        "style",
        "format",
    ]
    assert all(row["status"] == "completed" for row in stage_rows)
    assert source_count == 1

    markdown = format_output["markdown"]
    assert "# Transformers in Practice" in markdown
    assert "[cite:" not in markdown
    assert "[1]" in markdown
    assert "## References" in markdown
    assert "Attention Is All You Need" in markdown
    assert format_output["citations"] == 1
    assert format_output["render_status"] == "ok"

    assert document is not None
    assert document["byte_size"] > 1000
    assert document["content"][:2] == b"PK"  # a real DOCX archive


async def test_pipeline_resumes_after_a_completed_stage(db_pool, monkeypatch):
    """A second run must skip the stages that already succeeded."""
    from api.database import ensure_user
    from workers import stages
    from workers.worker import run_pipeline

    user_id = str(uuid.uuid4())
    job_id = str(uuid.uuid4())
    await ensure_user(user_id)

    async with db_pool.acquire() as conn:
        await conn.execute(
            """
            INSERT INTO jobs (id, user_id, topic, status, citation_style)
            VALUES ($1, $2, 'Resume behaviour', 'queued', 'apa')
            """,
            uuid.UUID(job_id),
            uuid.UUID(user_id),
        )
        await conn.execute(
            """
            INSERT INTO job_stages (id, job_id, stage_name, status, started_at,
                                    completed_at, output_data)
            VALUES ($1, $2, 'research', 'completed', now(), now(), $3)
            """,
            uuid.uuid4(),
            uuid.UUID(job_id),
            {"sources_found": 1},
        )

    async def fake_search(query, max_results=3):
        raise AssertionError("research stage should have been skipped")

    monkeypatch.setattr(stages, "get_search_results", fake_search)
    monkeypatch.setattr(stages.httpx, "AsyncClient", FakeHttpClient)

    async with db_pool.acquire() as conn:
        await conn.execute(
            """
            INSERT INTO sources (job_id, url, title, content_text, source_type)
            VALUES ($1, 'https://example.org/a', 'A source', $2, 'web')
            """,
            uuid.UUID(job_id),
            "Attention improved long sequence translation. " * 10,
        )

    await run_pipeline({"db_pool": db_pool, "provider_pool": FakeProviderPool()},
                       job_id, user_id)

    async with db_pool.acquire() as conn:
        stages_run = await conn.fetch(
            """
            SELECT stage_name FROM job_stages
            WHERE job_id = $1 AND stage_name <> 'research'
            """,
            uuid.UUID(job_id),
        )
        research_rows = await conn.fetchval(
            """
            SELECT count(*) FROM job_stages
            WHERE job_id = $1 AND stage_name = 'research'
            """,
            uuid.UUID(job_id),
        )

    assert {row["stage_name"] for row in stages_run} == {
        "outline",
        "draft",
        "verify",
        "style",
        "format",
    }
    assert research_rows == 1
